import io
import shutil
import tempfile
from copy import copy
from pathlib import Path

import openpyxl
from openpyxl.styles import Border, Side


# ── Border helper ─────────────────────────────────────────────────────────────
_THIN = Side(border_style="thin", color="000000")
_THIN_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def _apply_border(ws, min_row, max_row, min_col, max_col):
    for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
        for cell in row:
            cell.border = _THIN_BORDER


# ── Load data from uploaded files ────────────────────────────────────────────
def _load_data(uploaded_files):
    """Return merged gv_rows and hs_rows from all uploaded files."""
    gv_rows, hs_rows = [], []

    for uf in uploaded_files:
        wb = openpyxl.load_workbook(io.BytesIO(uf.read()), data_only=True)
        uf.seek(0)

        # GIAO-VIEN: sheet tên 'GIAO-VIEN', lấy A2:E
        if "GIAO-VIEN" in wb.sheetnames:
            ws = wb["GIAO-VIEN"]
            for row in ws.iter_rows(min_row=2, min_col=1, max_col=5, values_only=True):
                if any(v is not None and str(v).strip() != "" for v in row):
                    gv_rows.append(list(row))

        # HOC-SINH: sheet tên 'Danh sách HS toàn trường', lấy A7:I
        hs_sheet = None
        for name in wb.sheetnames:
            if "danh sách hs" in name.lower() or "danh sach hs" in name.lower():
                hs_sheet = wb[name]
                break
        if hs_sheet:
            for row in hs_sheet.iter_rows(min_row=7, min_col=1, max_col=9, values_only=True):
                if any(v is not None and str(v).strip() != "" for v in row):
                    hs_rows.append(list(row))

    return gv_rows, hs_rows


# ── Fill ADMIN sheet ──────────────────────────────────────────────────────────
def _fill_admin(ws, admin_info):
    """
    A1:D1 CỐ ĐỊNH, B2:B6 CỐ ĐỊNH
    A2=tên trường, C2=tk_admin, D2=mk_admin
    A3=ht.ten, C3=ht.tk, D3=ht.mk
    A4=hp1.ten, C4=hp1.tk, D4=hp1.mk
    A5=hp2.ten, C5=hp2.tk, D5=hp2.mk
    A6=hp3.ten, C6=hp3.tk, D6=hp3.mk
    """
    ws["A2"] = admin_info["ten_truong"] or None
    ws["C2"] = admin_info["tk_admin"] or None
    ws["D2"] = admin_info["mk_admin"] or None

    for i, row_data in enumerate(admin_info["rows"], start=3):
        ws[f"A{i}"] = row_data["ten"] or None
        ws[f"C{i}"] = row_data["tk"] or None
        ws[f"D{i}"] = row_data["mk"] or None


# ── Fill GIAO-VIEN sheet ─────────────────────────────────────────────────────
def _fill_giaovien(ws, gv_rows):
    if not gv_rows:
        return
    for i, row in enumerate(gv_rows, start=2):
        for j, val in enumerate(row, start=1):
            ws.cell(row=i, column=j, value=val)
    _apply_border(ws, min_row=2, max_row=1 + len(gv_rows), min_col=1, max_col=5)


# ── Fill HOC-SINH sheet ──────────────────────────────────────────────────────
def _fill_hochsinh(ws, hs_rows):
    if not hs_rows:
        return
    for i, row in enumerate(hs_rows, start=2):
        for j, val in enumerate(row, start=1):
            ws.cell(row=i, column=j, value=val)
    _apply_border(ws, min_row=2, max_row=1 + len(hs_rows), min_col=1, max_col=9)


# ── Save workbook to bytes ────────────────────────────────────────────────────
def _wb_to_bytes(wb):
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


# ── Build file name ───────────────────────────────────────────────────────────
def _build_name(template_name, ten_truong, suffix=""):
    base = Path(template_name).stem
    label = ten_truong or "Truong"
    name = f"{base} - {label}{suffix}.xlsx"
    return name


# ── Main export orchestrator ──────────────────────────────────────────────────
def process_export(uploaded_files, export_type, admin_info, templates_dir, ten_truong):
    gv_rows, hs_rows = _load_data(uploaded_files)
    label = ten_truong or ""
    result = []  # list of (filename, bytes)

    if export_type == "File Export gộp chung":
        tpl_path = templates_dir / "TK Onluyen (Admin, GV, HS).xlsx"
        wb = openpyxl.load_workbook(str(tpl_path))

        _fill_admin(wb["ADMIN"], admin_info)
        _fill_giaovien(wb["GIAO-VIEN"], gv_rows)
        _fill_hochsinh(wb["HOC-SINH"], hs_rows)

        fname = f"TK Onluyen (Admin, GV, HS) - {label}.xlsx" if label else "TK Onluyen (Admin, GV, HS).xlsx"
        result.append((fname, _wb_to_bytes(wb)))

    elif export_type == "File Export tách lẻ":
        tpl_path = templates_dir / "TK Onluyen (Admin, GV, HS).xlsx"
        tpl_wb = openpyxl.load_workbook(str(tpl_path))
        hdsd_ws = tpl_wb["HDSD"]

        # ── File ADMIN ────────────────────────────────────────────
        wb_admin = openpyxl.load_workbook(str(tpl_path))
        # Xoá sheet không cần, giữ ADMIN + HDSD
        for sname in wb_admin.sheetnames:
            if sname not in ("ADMIN", "HDSD"):
                del wb_admin[sname]
        _fill_admin(wb_admin["ADMIN"], admin_info)
        fname_admin = f"TK Onluyen - ADMIN - {label}.xlsx" if label else "TK Onluyen - ADMIN.xlsx"
        result.append((fname_admin, _wb_to_bytes(wb_admin)))

        # ── File GIAO-VIEN ────────────────────────────────────────
        wb_gv = openpyxl.load_workbook(str(tpl_path))
        for sname in wb_gv.sheetnames:
            if sname not in ("GIAO-VIEN", "HDSD"):
                del wb_gv[sname]
        _fill_giaovien(wb_gv["GIAO-VIEN"], gv_rows)
        fname_gv = f"TK Onluyen - GIAO-VIEN - {label}.xlsx" if label else "TK Onluyen - GIAO-VIEN.xlsx"
        result.append((fname_gv, _wb_to_bytes(wb_gv)))

        # ── File HOC-SINH ─────────────────────────────────────────
        wb_hs = openpyxl.load_workbook(str(tpl_path))
        for sname in wb_hs.sheetnames:
            if sname not in ("HOC-SINH", "HDSD"):
                del wb_hs[sname]
        _fill_hochsinh(wb_hs["HOC-SINH"], hs_rows)
        fname_hs = f"TK Onluyen - HOC-SINH - {label}.xlsx" if label else "TK Onluyen - HOC-SINH.xlsx"
        result.append((fname_hs, _wb_to_bytes(wb_hs)))

    elif export_type == "File Export Tiểu học":
        tpl_path = templates_dir / "TK Onluyen (Tiểu học).xlsx"
        wb = openpyxl.load_workbook(str(tpl_path))

        _fill_admin(wb["ADMIN"], admin_info)
        _fill_giaovien(wb["GIAO-VIEN"], gv_rows)
        _fill_hochsinh(wb["HOC-SINH"], hs_rows)

        fname = f"TK Onluyen (Tiểu học) - {label}.xlsx" if label else "TK Onluyen (Tiểu học).xlsx"
        result.append((fname, _wb_to_bytes(wb)))

    return result
